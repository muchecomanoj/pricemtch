import io
import json
import os
import random
import re
import shutil
import socket
import threading
import sys
import time
import uuid
import zipfile
from bs4 import BeautifulSoup
from flask import Flask, jsonify, render_template, request, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import requests
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium_stealth import stealth
import argparse
from marketplaces import MARKETPLACES, MARKETPLACE_DOMAINS, DEFAULT_DOMAIN
from openapi_spec import OPENAPI_SPEC

# --- ROBUST PYINSTALLER PATH BINDING ---
def get_resource_path(relative_path):
    """ Get absolute path to resource, working for dev and PyInstaller single-exe deployment """
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

app = Flask(__name__, template_folder=get_resource_path('templates'))

TEMP_DIR = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "temp_scraping_sessions")
os.makedirs(TEMP_DIR, exist_ok=True)

USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:109.0) Gecko/20100101 Firefox/121.0"
]

# The storefront registry lives in marketplaces.py so the Swagger spec shares it

def search_order(domain=None, domains=None):
    """ Marketplaces to try, in order: `domain` first, then either the caller's
    `domains` list or every marketplace. Duplicates are dropped. """
    pool = domains or MARKETPLACE_DOMAINS
    ordered = ([domain] if domain else []) + list(pool)
    return list(dict.fromkeys(ordered))

# The scraper's own fields, as opposed to the per-product technical details that
# vary listing by listing. Used to split the API response into the two groups.
CORE_FIELDS = ("ASIN", "Marketplace", "ProductName", "Description",
               "Feautured Product Description", "ImageURLs",
               "Price", "PriceValue", "Currency",
               "ListPrice", "ListPriceValue", "SavingsPercent", "PriceStatus",
               "OriginalPrice", "OriginalPriceValue", "OriginalCurrency",
               "FxRate", "FxAsOf", "FxSource")

def clean_amazon_image_url(url):
    if not url: return None
    return re.sub(r'\._[A-Z0-9_,\-]+_\.', '.', url)

# --- PRICE EXTRACTION ---
# Amazon ships several price layouts depending on marketplace, category and A/B
# bucket, so the price is hunted through these in order of how specific they are.
# The apex "accessibility label" is the one used on the newer INR/US pages; it
# reads e.g. "INR 1,662.28 with 32 percent savings".
# The product's OWN price containers. Everything below is searched strictly
# inside these: a detail page is full of sponsored tiles and recommendation
# carousels whose prices belong to completely different products.
PRICE_ROOTS = [
    "#corePriceDisplay_desktop_feature_div",
    "#corePrice_feature_div",
    "#corePrice_desktop",
    "#apex_desktop",
    "#desktop_unifiedPrice",
    "#unifiedPrice_feature_div",
]

PRICE_SELECTORS = [
    "span.priceToPay span.a-offscreen",
    "span.apex-price-to-pay span.a-offscreen",
    # On the apex layout priceToPay's a-offscreen is EMPTY and this label is the
    # only node carrying the real price, so it has to beat the generic selectors
    # below - otherwise they match the struck-through basis price instead.
    "#apex-pricetopay-accessibility-label",
    # :not(.a-text-price) keeps these off the struck-through list price
    "span.a-price:not(.a-text-price) span.a-offscreen",
]

# Older layouts that are their own leaf node rather than living under a root
LEGACY_PRICE_SELECTORS = [
    "#priceblock_ourprice",
    "#priceblock_dealprice",
    "#priceblock_saleprice",
    "#price_inside_buybox",
    "#newBuyBoxPrice",
    "#tp_price_block_total_price_ww span.a-offscreen",
]

# The struck-through "was"/RRP figure, again only within a price root
LIST_PRICE_SELECTORS = [
    "span.basisPrice span.a-offscreen",
    "span.apex-basisprice-value span.a-offscreen",
    "span.a-price.a-text-price span.a-offscreen",
]

# When every a-offscreen is empty, the visible price is still there as aria-hidden
# pieces (symbol / whole / fraction) and can be stitched back together.
COMPOSITE_PRICE_CONTAINERS = [
    "span.priceToPay",
    "span.apex-price-to-pay",
    "span.a-price:not(.a-text-price)",
]

SAVINGS_SELECTORS = ["span.savingsPercentage", "span.savingPriceOverride"]

# Each marketplace's home currency. A price that comes back in anything else is
# Amazon's localised "ship it to your country" price, not the domestic one.
CURRENCY_BY_DOMAIN = {domain: info["currency"] for domain, info in MARKETPLACES.items()}

ISO_CURRENCIES = sorted(set(CURRENCY_BY_DOMAIN.values()))
# Lookarounds rather than \b, so a code glued to digits ("INR4,990.67") still matches
CURRENCY_CODE_PATTERN = re.compile(r"(?<![A-Z])(" + "|".join(ISO_CURRENCIES) + r")(?![A-Z])")

# Checked in order, so multi-character symbols must precede the ones they
# contain: "CA$" before "A$", "US$" before "S$", "R$" before a bare "$".
CURRENCY_SYMBOLS = [
    ("CA$", "CAD"), ("AU$", "AUD"), ("MX$", "MXN"), ("US$", "USD"),
    ("R$", "BRL"), ("S$", "SGD"), ("A$", "AUD"), ("C$", "CAD"),
    ("zł", "PLN"), ("₺", "TRY"), ("£", "GBP"), ("€", "EUR"), ("₹", "INR"),
    ("￥", "JPY"), ("¥", "JPY"),
    ("د.إ", "AED"), ("ر.س", "SAR"), ("ج.م", "EGP"),
]
# Word-like markers that need boundaries so they don't fire inside other words
CURRENCY_WORD_MARKERS = [
    (re.compile(r"(?<![A-Za-z])kr(?![A-Za-z])"), "SEK"),
    (re.compile(r"(?<![A-Za-z])TL(?![A-Za-z])"), "TRY"),
    # South Africa writes "R 299.00"; a lone R directly before digits
    (re.compile(r"(?<![A-Za-z])R\s?(?=\d)"), "ZAR"),
]
# Storefronts whose currency is written with a bare "$"
DOLLAR_CURRENCIES = {"USD", "CAD", "AUD", "SGD", "MXN"}
# A number possibly carrying thousands/decimal separators: 1,662.28 / 1.662,28 /
# 1 299,00. Grouping separators must be followed by exactly three digits, so a
# price can't fuse with a neighbouring number ("479.13 ... 14%").
PRICE_NUMBER_PATTERN = re.compile(r"\d+(?:[.,\s ]\d{3})*(?:[.,]\d{1,2})?")
PERCENT_PATTERN = re.compile(r"(\d+(?:[.,]\d+)?)\s*(?:%|percent)", re.IGNORECASE)

# A currency marker written after the number: "17,50 €", "129 kr", "49,99 zł"
SUFFIX_MARKER_PATTERN = re.compile(
    r"\s?(?:" + "|".join(re.escape(symbol) for symbol, _ in CURRENCY_SYMBOLS)
    + r"|kr|TL|" + "|".join(ISO_CURRENCIES) + r")(?![A-Za-z])"
)

def clean_price_text(text):
    """ Collapse whitespace and keep only the price itself. The apex label
    appends a savings clause in the storefront's language - "with 32 percent
    savings", "con un ahorro del 14%", "avec 10 % d'economies" - so rather than
    split on any particular word, cut right after the first number (plus a
    trailing currency marker, for storefronts that write it after). """
    if not text:
        return None
    cleaned = " ".join(text.replace(" ", " ").split())
    number = PRICE_NUMBER_PATTERN.search(cleaned)
    if not number:
        return cleaned or None
    end = number.end()
    suffix = SUFFIX_MARKER_PATTERN.match(cleaned, end)
    if suffix:
        end = suffix.end()
    return cleaned[:end].strip() or None

def parse_price_amount(text):
    """ Pull a float out of a price string, honouring both 1,662.28 and 1.662,28 """
    if not text:
        return None
    match = PRICE_NUMBER_PATTERN.search(text.replace(" ", " "))
    if not match:
        return None

    token = re.sub(r"\s", "", match.group(0)).strip(".,")
    separator_at = max(token.rfind("."), token.rfind(","))
    try:
        if separator_at == -1:
            return float(token)
        tail = token[separator_at + 1:]
        # 1-2 trailing digits means it was the decimal separator; 3 means thousands
        if len(tail) in (1, 2) and tail.isdigit():
            whole = re.sub(r"[.,]", "", token[:separator_at]) or "0"
            return float(f"{whole}.{tail}")
        return float(re.sub(r"[.,]", "", token))
    except ValueError:
        return None

def currency_in_text(text, domain=None):
    """ The currency a price string is written in, or None if it carries no
    marker at all. Explicit ISO code wins, then symbols, then word markers. A
    bare "$" is resolved by the marketplace (.ca -> CAD, .com.au -> AUD, ...). """
    if not text:
        return None
    code = CURRENCY_CODE_PATTERN.search(text.upper())
    if code:
        return code.group(1)
    for symbol, currency in CURRENCY_SYMBOLS:
        if symbol in text:
            return currency
    for pattern, currency in CURRENCY_WORD_MARKERS:
        if pattern.search(text):
            return currency
    if "$" in text:
        home = CURRENCY_BY_DOMAIN.get(domain)
        return home if home in DOLLAR_CURRENCIES else "USD"
    return None

def detect_currency(text, domain):
    return currency_in_text(text, domain) or CURRENCY_BY_DOMAIN.get(domain)

def looks_like_price(text):
    """ Text only counts as a price if it carries a currency marker AND a digit,
    wherever the marker sits ("£12.99", "17,50 €", "129 kr"). Guards against
    a-offscreen nodes that hold titles or rating text - common in buy-box and
    carousel markup. """
    if not text or len(text) > 40:
        return False
    return any(c.isdigit() for c in text) and currency_in_text(text) is not None

def select_first_price_text(root, selectors):
    for selector in selectors:
        for node in root.select(selector):
            text = clean_price_text(node.get_text())
            if looks_like_price(text):
                return text
    return None

def price_roots(soup):
    """ The product's own price containers, in priority order """
    return [node for selector in PRICE_ROOTS for node in soup.select(selector)]

def compose_price_from_parts(root):
    """ Rebuild the price from its aria-hidden pieces when no a-offscreen carries
    it. The pieces must be reassembled explicitly - reading the rendered node's
    text can silently drop the decimal point and inflate the price 100x. """
    for selector in COMPOSITE_PRICE_CONTAINERS:
        for node in root.select(selector):
            whole = node.select_one("span.a-price-whole")
            if not whole:
                continue
            whole_text = "".join(whole.get_text().split()).strip(".,")
            if not whole_text.replace(",", "").replace(".", "").isdigit():
                continue
            symbol = node.select_one("span.a-price-symbol")
            fraction = node.select_one("span.a-price-fraction")
            text = (symbol.get_text(strip=True) if symbol else "") + whole_text
            if fraction:
                fraction_text = "".join(fraction.get_text().split())
                if fraction_text.isdigit():
                    text += "." + fraction_text
            if looks_like_price(text):
                return text
    return None

def extract_price(soup, domain):
    """ Returns the price block, or None when the listing shows no price of its
    own (out of stock, "see all buying options", can't ship here, ...).

    Everything is read from inside the product's own price containers, so a
    sponsored tile elsewhere on the page can never be mistaken for the price.
    """
    raw = list_raw = savings_text = None
    for root in price_roots(soup):
        raw = select_first_price_text(root, PRICE_SELECTORS) or compose_price_from_parts(root)
        if raw:
            list_raw = select_first_price_text(root, LIST_PRICE_SELECTORS)
            savings_text = select_first_price_text(root, SAVINGS_SELECTORS)
            if not savings_text:
                label = root.select_one("#apex-pricetopay-accessibility-label")
                savings_text = " ".join(label.get_text().split()) if label else None
            break

    if not raw:
        for selector in LEGACY_PRICE_SELECTORS:
            node = soup.select_one(selector)
            text = clean_price_text(node.get_text()) if node else None
            if looks_like_price(text):
                raw = text
                break

    amount = parse_price_amount(raw)
    if amount is None:
        return None

    list_amount = parse_price_amount(list_raw)
    # A "list price" at or below the price being paid is a mis-grab (usually the
    # current price echoed back), so don't report it as an RRP.
    if list_amount is not None and list_amount <= amount:
        list_raw, list_amount = None, None

    savings_percent = None
    if savings_text:
        percent = PERCENT_PATTERN.search(savings_text)
        if percent:
            savings_percent = parse_price_amount(percent.group(1))
    # Derive it when Amazon didn't print a savings badge
    if savings_percent is None and list_amount and list_amount > amount:
        savings_percent = round((list_amount - amount) / list_amount * 100)

    currency = detect_currency(raw, domain)
    return {
        "raw": raw,
        "amount": amount,
        "currency": currency,
        "list_price_raw": list_raw if list_amount is not None else None,
        "list_price_amount": list_amount,
        "savings_percent": savings_percent,
        # False means Amazon served this storefront's "ship it abroad" pricing
        # instead of its domestic pricing - see PRICE_STATUS_LOCALISED.
        "is_marketplace_currency": currency == CURRENCY_BY_DOMAIN.get(domain),
    }

# Any 10-character alphanumeric token: covers B0-style ASINs, plain ISBN-10s
# like 5585222953, and the X-suffixed ISBN variant.
ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$", re.IGNORECASE)

def normalize_cell_to_asin(cell):
    """ Excel stores an all-digit ASIN as a number, so coerce those back to a 10-char string """
    if cell is None:
        return None
    if isinstance(cell, str):
        return cell.strip()
    if isinstance(cell, int):
        return str(cell)
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return None

@app.route("/")
def home():
    return render_template("index.html", marketplaces=MARKETPLACES, default_domain=DEFAULT_DOMAIN)

@app.route("/docs")
def api_docs():
    """ Swagger UI, reading the spec from /openapi.json """
    return render_template("swagger.html")

@app.route("/openapi.json")
def openapi_json():
    return jsonify(OPENAPI_SPEC)

@app.route("/upload-excel", methods=["POST"])
def upload_excel():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400

    try:
        in_memory_file = io.BytesIO(file.read())
        workbook = openpyxl.load_workbook(in_memory_file, data_only=True)
        sheet = workbook.active

        asins = []
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                clean_cell = normalize_cell_to_asin(cell)
                if clean_cell and ASIN_PATTERN.match(clean_cell):
                    asins.append(clean_cell.upper())
        asins = list(dict.fromkeys(asins))
    except Exception as e:
        return jsonify({"error": f"Failed to parse Excel schema: {str(e)}"}), 400

    if not asins:
        return jsonify({"error": "No valid Amazon ASINs found in the active sheet."}), 400

    session_id = str(uuid.uuid4())
    os.makedirs(os.path.join(TEMP_DIR, session_id), exist_ok=True)

    return jsonify({"asins": asins, "session_id": session_id})

# Price status reported alongside every scrape
PRICE_STATUS_OK = "ok"
PRICE_STATUS_CONVERTED = "converted"
PRICE_STATUS_LOCALISED = "localised_currency"
PRICE_STATUS_UNAVAILABLE = "unavailable"

# --- CURRENCY CONVERSION ---
# Amazon prices by the visitor's IP, so a storefront abroad quotes in the
# scraper's own currency. Rather than withholding that number, it is converted
# into the marketplace's currency. The original is always kept alongside.
FX_PRIMARY_URL = "https://open.er-api.com/v6/latest/{base}"
# ECB data: no AED, SAR or EGP, so those only convert via the primary
FX_FALLBACK_URL = "https://api.frankfurter.app/latest?from={base}"
FX_TTL_SECONDS = int(os.environ.get("AMAZON_FX_TTL", 6 * 3600))
# Currencies without a distinctive symbol are written with their ISO code
SYMBOL_BY_CURRENCY = {"GBP": "£", "USD": "$", "EUR": "€", "INR": "₹", "CAD": "CA$",
                      "AUD": "A$", "SGD": "S$", "MXN": "MX$", "BRL": "R$", "ZAR": "R",
                      "JPY": "¥", "TRY": "₺"}
# Currencies with no minor unit
ZERO_DECIMAL_CURRENCIES = {"JPY"}

def round_money(amount, currency):
    return round(amount, 0 if currency in ZERO_DECIMAL_CURRENCIES else 2)

_fx_cache = {}
_fx_lock = threading.Lock()

def fetch_fx_rates(base):
    """ Rates FROM `base` to every other currency, cached for FX_TTL_SECONDS.
    Returns None when both providers are unreachable. """
    now = time.time()
    with _fx_lock:
        cached = _fx_cache.get(base)
        if cached and now - cached["fetched_at"] < FX_TTL_SECONDS:
            return cached

    payload = None
    try:
        data = requests.get(FX_PRIMARY_URL.format(base=base), timeout=12).json()
        if data.get("rates"):
            payload = {"rates": data["rates"],
                       "as_of": data.get("time_last_update_utc"),
                       "source": "open.er-api.com"}
    except Exception:
        payload = None

    if payload is None:
        try:
            data = requests.get(FX_FALLBACK_URL.format(base=base), timeout=12).json()
            if data.get("rates"):
                payload = {"rates": data["rates"],
                           "as_of": data.get("date"),
                           "source": "frankfurter.app"}
        except Exception:
            payload = None

    if payload is None:
        return None

    payload["fetched_at"] = now
    with _fx_lock:
        _fx_cache[base] = payload
    return payload

def convert_amount(amount, from_currency, to_currency):
    """ -> (converted_amount, rate, as_of, source), or None if no rate is available """
    if amount is None:
        return None
    if from_currency == to_currency:
        return amount, 1.0, None, None
    rates = fetch_fx_rates(from_currency)
    if not rates:
        return None
    rate = rates["rates"].get(to_currency)
    if not rate:
        return None
    return round_money(amount * rate, to_currency), rate, rates.get("as_of"), rates.get("source")

def format_price(amount, currency):
    decimals = 0 if currency in ZERO_DECIMAL_CURRENCIES else 2
    number = f"{amount:,.{decimals}f}"
    symbol = SYMBOL_BY_CURRENCY.get(currency)
    return f"{symbol}{number}" if symbol else f"{currency} {number}"

# Every price column, so a row never carries a stale value from another branch
PRICE_FIELDS = ("Price", "PriceValue", "Currency", "ListPrice", "ListPriceValue",
                "SavingsPercent", "OriginalPrice", "OriginalPriceValue",
                "OriginalCurrency", "FxRate", "FxAsOf", "FxSource")

def blank_price_fields(features):
    for field in PRICE_FIELDS:
        features[field] = ""

def proxy_for_domain(domain):
    """ Amazon prices by the visitor's IP, so the ONLY way to read a storefront's
    domestic prices is to reach it from that country. Point a marketplace at a
    proxy located there via an environment variable:

        AMAZON_PROXY_CO_UK=http://user:pass@london.proxy:8000
        AMAZON_PROXY_COM=http://user:pass@us.proxy:8000
        AMAZON_PROXY=http://fallback.proxy:8000      (used for any marketplace)

    Without one, a storefront outside your own country returns its "ship it to
    your country" pricing, which is a different number in a different currency.
    """
    specific = "AMAZON_PROXY_" + domain.upper().replace(".", "_").replace("-", "_")
    return os.environ.get(specific) or os.environ.get("AMAZON_PROXY") or None

class DriverLaunchError(RuntimeError):
    """ Chrome itself could not be started (browser missing, driver mismatch, ...) """

def build_driver(proxy=None):
    chrome_options = Options()
    if proxy:
        chrome_options.add_argument(f"--proxy-server={proxy}")
    chrome_options.add_argument("--headless=new")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--no-sandbox")
    chrome_options.add_argument("--disable-dev-shm-usage")
    chrome_options.add_argument("--window-size=1920,1080")
    chrome_options.add_argument("--disable-blink-features=AutomationControlled")
    chrome_options.add_experimental_option("excludeSwitches", ["enable-automation"])
    chrome_options.add_experimental_option('useAutomationExtension', False)

    try:
        driver = webdriver.Chrome(options=chrome_options)
        driver.set_page_load_timeout(PAGE_LOAD_TIMEOUT)
        stealth(driver, languages=["en-US", "en"], vendor="Google Inc.", platform="Win32", webgl_vendor="Intel Inc.", renderer="Intel Iris OpenGL Engine", fix_hairline=True)
        driver.execute_cdp_cmd("Network.setUserAgentOverride", {"userAgent": random.choice(USER_AGENTS)})
    except Exception as driver_err:
        raise DriverLaunchError(str(driver_err))
    return driver

# --- PER-MARKETPLACE PAGE OUTCOMES ---
PAGE_FOUND = "found"
PAGE_NOT_FOUND = "not_found"            # Amazon's own 404 for this ASIN
PAGE_BLOCKED = "blocked"                # CAPTCHA or "Server Busy" throttle page
PAGE_NO_PRODUCT = "no_product_page"     # a full page loaded, but not a product
PAGE_ERROR = "error"                    # navigation failed or timed out

PAGE_LOAD_TIMEOUT = 30
# Real pages are 1-2 MB; 404, CAPTCHA and throttle stubs are all under ~4 KB
STUB_PAGE_BYTES = 20000

# Chrome exposes the document's HTTP status through Navigation Timing
HTTP_STATUS_JS = ("var n = performance.getEntriesByType('navigation')[0];"
                  "return n ? n.responseStatus : null;")

def http_status(driver):
    try:
        status = driver.execute_script(HTTP_STATUS_JS)
        return int(status) if status else None
    except Exception:
        return None

def classify_page(html, status=None):
    """ Decide what a marketplace answered, without a full parse.

    Every signal is language-independent. The HTTP status is the primary one:
    404 pages are titled "Page Not Found", "Documento no encontrado", "Nie
    znaleziono strony"... and the Americas storefronts don't carry the cs_404
    link the others do, but all of them answer HTTP 404. The CAPTCHA form always
    posts to validateCaptcha. """
    if 'id="productTitle"' in html:
        return PAGE_FOUND
    if "validateCaptcha" in html:
        return PAGE_BLOCKED
    if status == 404 or "cs_404" in html:
        return PAGE_NOT_FOUND
    if status in (429, 503) or len(html) < STUB_PAGE_BYTES:
        return PAGE_BLOCKED  # "Server Busy" throttling and similar stubs
    return PAGE_NO_PRODUCT

def open_product_page(driver, domain, asin):
    """ Load one marketplace's product page. Returns (outcome, attempts).

    driver.get() already blocks until the page's load event, so only a short
    settle is added for found pages (tested: 1.2s captures exactly what the old
    fixed 4-5.5s wait did). A blocked page is retried once after a back-off,
    because "Server Busy" throttling routinely clears on the second request -
    it must never be recorded as the ASIN being absent. """
    url = f"https://www.amazon.{domain}/dp/{asin}/"
    outcome = PAGE_ERROR
    attempts = 0
    for attempts in (1, 2):
        try:
            driver.get(url)
            outcome = classify_page(driver.page_source, http_status(driver))
        except Exception:
            outcome = PAGE_ERROR
        if outcome not in (PAGE_BLOCKED, PAGE_ERROR):
            break
        if attempts == 1:
            time.sleep(random.uniform(2.0, 3.5))

    if outcome == PAGE_FOUND:
        time.sleep(random.uniform(1.0, 2.0))
    else:
        time.sleep(random.uniform(0.3, 0.8))  # pace consecutive storefronts
    return outcome, attempts

# A price counts only if it is in the marketplace's currency, natively or converted
PRICED_STATUSES = (PRICE_STATUS_OK, PRICE_STATUS_CONVERTED)
PAGE_FOUND_NO_PRICE = "found_no_price"  # listed here, but no usable price

def parse_product_page(soup, asin, found_domain):
    """ Extract everything from one storefront's product page.
    Returns (features, image_urls). """
    # Create a dictionary to hold features for this specific ASIN execution
    asin_features = {"ASIN": asin, "Marketplace": found_domain}

    # 1. Title Extraction
    title_span = soup.find("span", id="productTitle")
    asin_features["ProductName"] = title_span.get_text(strip=True) if title_span else "Not Found"

    # 2. Description Parsing
    description_text = "Not Found"
    about_heading = soup.find('h1', string=re.compile(r'about this item', re.IGNORECASE))
    if about_heading:
        next_node = about_heading.find_next_sibling()
        while next_node and next_node.name != 'ul':
            internal_ul = next_node.find('ul')
            if internal_ul: next_node = internal_ul; break
            next_node = next_node.find_next_sibling()
        if not next_node or next_node.name != 'ul': next_node = about_heading.find('ul')
        if next_node and next_node.name == 'ul':
            bullets = [li.get_text(strip=True) for li in next_node.find_all('li') if li.get_text(strip=True)]
            description_text = "\n".join(bullets)
    if description_text == "Not Found":
        # The heading is localised ("Über diesen Artikel", "この商品について", ...)
        # but the bullet container's id is the same on every storefront
        container = soup.select_one("#feature-bullets") or soup.select_one("#featurebullets_feature_div")
        if container:
            bullets = [li.get_text(strip=True) for li in container.select("ul li") if li.get_text(strip=True)]
            if bullets:
                description_text = "\n".join(bullets)
    asin_features["Description"] = description_text

    # 2b. Featured product description (only some listings carry this block)
    featured_description = ""
    product_desc_div = soup.find("div", id="productDescription")
    if product_desc_div:
        featured_parts = []
        for para in product_desc_div.find_all("p"):
            for span in para.find_all("span"):
                text = span.get_text(strip=True)
                if text and text not in featured_parts:
                    featured_parts.append(text)
        featured_description = "\n".join(featured_parts)
    asin_features["Feautured Product Description"] = featured_description

    # 2c. Price, always reported in the marketplace's own currency. When the
    # storefront quotes in a foreign currency (it geolocated the scraper
    # abroad) the figure is converted at the day's rate, and the original is
    # kept so the conversion can always be audited.
    price_info = extract_price(soup, found_domain)
    target_currency = CURRENCY_BY_DOMAIN.get(found_domain)
    conversion = None

    if not price_info:
        price_status = PRICE_STATUS_UNAVAILABLE
    elif price_info["is_marketplace_currency"]:
        price_status = PRICE_STATUS_OK
    else:
        conversion = convert_amount(price_info["amount"], price_info["currency"], target_currency)
        price_status = PRICE_STATUS_CONVERTED if conversion else PRICE_STATUS_LOCALISED

    asin_features["PriceStatus"] = price_status
    blank_price_fields(asin_features)

    if price_status == PRICE_STATUS_OK:
        asin_features["Price"] = price_info["raw"]
        asin_features["PriceValue"] = price_info["amount"]
        asin_features["Currency"] = price_info["currency"]
        asin_features["ListPrice"] = price_info["list_price_raw"] or ""
        asin_features["ListPriceValue"] = price_info["list_price_amount"] or ""
        asin_features["SavingsPercent"] = price_info["savings_percent"] or ""

    elif price_status == PRICE_STATUS_CONVERTED:
        amount, rate, as_of, source = conversion
        asin_features["Price"] = format_price(amount, target_currency)
        asin_features["PriceValue"] = amount
        asin_features["Currency"] = target_currency
        asin_features["SavingsPercent"] = price_info["savings_percent"] or ""
        # The list price rides the same rate, so the discount stays consistent
        if price_info["list_price_amount"] is not None:
            list_amount = round(price_info["list_price_amount"] * rate, 2)
            asin_features["ListPrice"] = format_price(list_amount, target_currency)
            asin_features["ListPriceValue"] = list_amount
        asin_features["OriginalPrice"] = price_info["raw"]
        asin_features["OriginalPriceValue"] = price_info["amount"]
        asin_features["OriginalCurrency"] = price_info["currency"]
        asin_features["FxRate"] = rate
        asin_features["FxAsOf"] = as_of or ""
        asin_features["FxSource"] = source or ""

    elif price_status == PRICE_STATUS_LOCALISED:
        # Foreign currency AND no exchange rate available - report what we saw
        asin_features["OriginalPrice"] = price_info["raw"]
        asin_features["OriginalPriceValue"] = price_info["amount"]
        asin_features["OriginalCurrency"] = price_info["currency"]

    # 3. Dynamic Technical Detail Tables
    prod_details_div = soup.find("div", id="prodDetails")
  
    if prod_details_div:
        for table in prod_details_div.find_all("table", class_="prodDetTable"):
            for row in table.find_all("tr"):
                th, td = row.find("th"), row.find("td")
                if th and td:
                    key = th.get_text(strip=True)
                    val = " ".join(td.get_text().split())
                    # Standardize keys to avoid minor duplication bugs
                    if key:
                        asin_features[key] = val
    

    # 3b. Detail bullets: the "Key : Value" list that books everywhere, and
    # whole storefronts such as .co.jp, use instead of the prodDetails tables.
    # Table rows win if a page carries both.
    for li in soup.select("#detailBullets_feature_div li"):
        text = " ".join(li.get_text().replace("‏", "").replace("‎", "").split())
        key, sep, val = text.partition(":")
        key, val = key.strip(), val.strip()
        if sep and key and val and len(key) <= 60 and key not in asin_features:
            asin_features[key] = val

    # 4. Extract Images Links (no download — just capture the URLs)
    image_urls = []
    img_wrappers = soup.find_all("div", class_="imgTagWrapper")
    for wrapper in img_wrappers:
        img_tag = wrapper.find("img")
        if img_tag:
            for attr in ["data-old-hires", "src", "data-a-dynamic-image"]:
                if img_tag.has_attr(attr) and img_tag[attr]:
                    val = img_tag[attr]
                    if val.startswith("{"):
                        try:
                            keys = list(json.loads(val).keys())
                            if keys: image_urls.append(keys[0])
                        except Exception: pass
                    elif not val.startswith("data:image"): image_urls.append(val)

    alt_images_div = soup.find("div", id="altImages")
    if alt_images_div:
        for img in alt_images_div.find_all("img"):
            if img.has_attr("src") and not img["src"].startswith("data:image"): image_urls.append(img["src"])

    final_cleaned_urls = []
    for raw_url in image_urls:
        cleaned = clean_amazon_image_url(raw_url)
        if cleaned and cleaned not in final_cleaned_urls:
            if "play-button" not in cleaned.lower() and "icon" not in cleaned.lower(): final_cleaned_urls.append(cleaned)

    asin_features["ImageURLs"] = ", ".join(final_cleaned_urls) if final_cleaned_urls else ""

    return asin_features, final_cleaned_urls

def scrape_asin(asin, domain=DEFAULT_DOMAIN, domains=None, require_price=False):
    """ Scrape one ASIN, trying `domain` first and then every other marketplace
    (or just `domains`, if given) until the product page turns up.

    With require_price, a storefront that lists the product but shows no usable
    price is logged as "found_no_price" and the search carries on; the first
    storefront WITH a price wins. If none has one, the first listing found is
    returned rather than nothing - the product data is still worth having.

    Returns (features, image_urls, search_log). `features` is the flat field map
    that also feeds the Excel columns; its "Marketplace" is "Not Found" when the
    ASIN wasn't found. `search_log` records what each marketplace answered.
    Raises DriverLaunchError if Chrome won't start.
    """
    domains_to_try = search_order(domain, domains)
    search_log = []
    driver = None
    active_proxy = False  # sentinel: no driver built yet
    unpriced_listing = None  # first (features, images) found without a price

    try:
        for domain_attempt in domains_to_try:
            # Each marketplace may route through its own country's proxy, so the
            # driver is rebuilt whenever the proxy changes between attempts.
            proxy = proxy_for_domain(domain_attempt)
            if driver is None or proxy != active_proxy:
                if driver is not None:
                    driver.quit()
                driver = build_driver(proxy)
                active_proxy = proxy

            started = time.time()
            outcome, attempts = open_product_page(driver, domain_attempt, asin)
            entry = {"marketplace": domain_attempt, "result": outcome,
                     "attempts": attempts, "seconds": round(time.time() - started, 1)}
            search_log.append(entry)
            if outcome != PAGE_FOUND:
                continue

            features, images = parse_product_page(
                BeautifulSoup(driver.page_source, "lxml"), asin, domain_attempt)
            if not require_price or features.get("PriceStatus") in PRICED_STATUSES:
                return features, images, search_log

            entry["result"] = PAGE_FOUND_NO_PRICE
            if unpriced_listing is None:
                unpriced_listing = (features, images)

        if unpriced_listing is not None:
            return unpriced_listing[0], unpriced_listing[1], search_log

        # ASIN wasn't found on any marketplace we checked
        return ({"ASIN": asin, "ProductName": "Not Found", "Description": "Not Found", "Marketplace": "Not Found"},
                [], search_log)
    finally:
        if driver is not None:
            driver.quit()

@app.route("/scrape-single-asin", methods=["POST"])
def scrape_single_asin():
    data = request.json or {}
    asin = data.get("asin")
    session_id = data.get("session_id")
    domain = data.get("domain", DEFAULT_DOMAIN)

    if not asin or not session_id:
        return jsonify({"error": "Missing execution context parameters"}), 400
    # It's interpolated into the URL, so only known storefronts get through
    if domain not in MARKETPLACES:
        return jsonify({"error": f"Unsupported domain '{domain}'"}), 400

    session_folder = os.path.join(TEMP_DIR, session_id, asin)
    os.makedirs(session_folder, exist_ok=True)

    try:
        asin_features, _images, _search_log = scrape_asin(asin, domain)
    except DriverLaunchError as driver_err:
        return jsonify({"error": f"Failed to launch Chrome background driver: {str(driver_err)}"}), 500
    except Exception as parse_err:
        return jsonify({"error": f"Internal parsing routine failure: {str(parse_err)}"}), 500

    # Save data to a temporary JSON file for this ASIN so we can build a master file later
    with open(os.path.join(session_folder, "data.json"), "w", encoding="utf-8") as f:
        json.dump(asin_features, f, ensure_ascii=False, indent=4)

    return jsonify({"success": True})

def build_price_response(features):
    """ The price in the marketplace's own currency, carrying its provenance when
    it came from a conversion. None when the page showed no usable price. """
    if not features.get("Price"):
        return None

    price = {
        "raw": features.get("Price"),
        "amount": features.get("PriceValue"),
        "currency": features.get("Currency"),
        "list_price_raw": features.get("ListPrice") or None,
        "list_price_amount": features.get("ListPriceValue") or None,
        "savings_percent": features.get("SavingsPercent") if features.get("SavingsPercent") != "" else None,
        "converted": features.get("PriceStatus") == PRICE_STATUS_CONVERTED,
    }

    if price["converted"]:
        price["original"] = {
            "raw": features.get("OriginalPrice"),
            "amount": features.get("OriginalPriceValue"),
            "currency": features.get("OriginalCurrency"),
        }
        price["fx"] = {
            "rate": features.get("FxRate"),
            "as_of": features.get("FxAsOf") or None,
            "source": features.get("FxSource") or None,
        }
        price["note"] = (
            f"Converted from {features.get('OriginalCurrency')} at "
            f"{features.get('FxRate')}. amazon.{features.get('Marketplace')} quoted this listing in "
            f"{features.get('OriginalCurrency')} because it geolocated the scraper outside its "
            f"country, so the original includes import duty and international shipping. Treat this "
            f"as an estimate in {features.get('Currency')}, not the domestic shelf price."
        )
    return price

@app.route("/api/product", methods=["POST"])
def api_product():
    """ Stateless single-ASIN lookup — no upload, no session, no zip.

        POST /api/product   {"asin": "B0016URDD0"}
        POST /api/product   {"asin": "B0016URDD0", "domain": "de"}
        POST /api/product   {"asin": "B0016URDD0", "domains": ["fr", "it", "es"]}

    "domain" decides which marketplace is tried FIRST; every other storefront
    is then searched until the ASIN turns up. "domains" narrows the search to
    just those storefronts, in that order. 200 with the product, 404 if no
    searched marketplace carried the ASIN.
    """
    data = request.get_json(silent=True) or {}

    asin = normalize_cell_to_asin(data.get("asin"))
    if not asin:
        return jsonify({"error": "Missing 'asin' in request body"}), 400

    asin = asin.upper()
    if not ASIN_PATTERN.match(asin):
        return jsonify({"error": f"'{asin}' is not a valid 10-character ASIN"}), 400

    domains = data.get("domains")
    if domains is not None:
        if not isinstance(domains, list) or not domains or not all(isinstance(d, str) for d in domains):
            return jsonify({"error": "'domains' must be a non-empty list of marketplace domains"}), 400
        unknown = [d for d in domains if d not in MARKETPLACES]
        if unknown:
            return jsonify({"error": f"Unsupported domain(s) {unknown}. See GET /api/marketplaces"}), 400

    # With an explicit list and no starting point, the list's own order rules
    domain = data.get("domain") or (None if domains else DEFAULT_DOMAIN)
    if domain is not None and domain not in MARKETPLACES:
        return jsonify({"error": f"Unsupported domain '{domain}'. See GET /api/marketplaces"}), 400

    require_price = data.get("require_price", False)
    if not isinstance(require_price, bool):
        return jsonify({"error": "'require_price' must be true or false"}), 400

    try:
        features, images, search_log = scrape_asin(asin, domain, domains, require_price)
    except DriverLaunchError as driver_err:
        return jsonify({"error": f"Failed to launch Chrome background driver: {str(driver_err)}"}), 500
    except Exception as parse_err:
        return jsonify({"error": f"Internal parsing routine failure: {str(parse_err)}"}), 500

    if features.get("Marketplace") == "Not Found":
        # A storefront that throttled us or errored didn't actually say "no" - so
        # the miss is only conclusive if every storefront gave a real answer.
        inconclusive = [entry["marketplace"] for entry in search_log
                        if entry["result"] in (PAGE_BLOCKED, PAGE_ERROR)]
        return jsonify({
            "found": False,
            "asin": asin,
            "marketplaces_tried": [entry["marketplace"] for entry in search_log],
            "conclusive": not inconclusive,
            "unchecked_marketplaces": inconclusive,
            "search_log": search_log,
            "error": ("ASIN not found on any searched marketplace" if not inconclusive else
                      f"ASIN not found, but {len(inconclusive)} marketplace(s) blocked or failed the "
                      f"request, so it may exist there - retry with \"domains\": {inconclusive}")
        }), 404

    response = {
        "found": True,
        "asin": asin,
        "marketplace": features.get("Marketplace"),
        "product_name": features.get("ProductName"),
        "description": features.get("Description"),
        "featured_description": features.get("Feautured Product Description", ""),
        "price": build_price_response(features),
        "price_status": features.get("PriceStatus", PRICE_STATUS_UNAVAILABLE),
        "images": images,
        # Whatever the listing's technical-detail tables carried — brand, model,
        # dimensions, rank, reviews, ... varies product by product.
        "details": {k: v for k, v in features.items() if k not in CORE_FIELDS},
        # What each storefront answered on the way to finding it
        "search_log": search_log,
    }

    # Asked for a priced storefront, but every listing lacked a price: say so,
    # rather than let the first listing pass for a match
    if require_price and features.get("PriceStatus") not in PRICED_STATUSES:
        listed = [e["marketplace"] for e in search_log if e["result"] == PAGE_FOUND_NO_PRICE]
        response["note"] = (
            f"No searched storefront showed a price. The ASIN is listed on {', '.join(listed)} "
            f"without one; returning the first listing, amazon.{features.get('Marketplace')}."
        )

    # Foreign currency AND no exchange rate reachable, so nothing could be converted
    if features.get("PriceStatus") == PRICE_STATUS_LOCALISED:
        response["price_localised"] = {
            "raw": features.get("OriginalPrice"),
            "currency": features.get("OriginalCurrency"),
            "note": (
                f"amazon.{features.get('Marketplace')} quoted this listing in "
                f"{features.get('OriginalCurrency')}, and no exchange rate could be fetched to "
                f"convert it into {CURRENCY_BY_DOMAIN.get(features.get('Marketplace'))}. Check the "
                f"server's internet access and retry."
            ),
        }

    # Amazon sometimes serves a parent/variant listing whose detail table reports a
    # different ASIN than the one requested; surface it instead of silently swapping.
    scraped_asin = features.get("ASIN")
    if scraped_asin and scraped_asin != asin:
        response["listing_asin"] = scraped_asin

    return jsonify(response)

@app.route("/api/marketplaces", methods=["GET"])
def api_marketplaces():
    """ Every storefront the scraper can search, in default fallback order """
    return jsonify({
        "default": DEFAULT_DOMAIN,
        "marketplaces": [
            {"domain": domain, "url": f"https://www.amazon.{domain}",
             "country": info["country"], "currency": info["currency"]}
            for domain, info in MARKETPLACES.items()
        ],
    })

@app.route("/download-zip", methods=["GET"])
def download_zip():
    session_id = request.args.get("session_id")
    if not session_id: return "Error: Missing Session identifier", 400

    target_session_path = os.path.join(TEMP_DIR, session_id)
    if not os.path.exists(target_session_path): return "Error: Dynamic data session context expired.", 404

    # --- NEW ARCHITECTURE: COMPILE DYNAMIC MASTER EXCEL FILE ---
    all_asin_data = []
    global_headers = ["ASIN", "Marketplace", "ProductName", "Description", "Feautured Product Description"] # Core tracking blueprint order

    # 1. Look through all scraped ASIN folders to extract features map
    if os.path.exists(target_session_path):
        for item in os.listdir(target_session_path):
            item_path = os.path.join(target_session_path, item)
            if os.path.isdir(item_path):
                json_file_path = os.path.join(item_path, "data.json")
                if os.path.exists(json_file_path):
                    with open(json_file_path, "r", encoding="utf-8") as f:
                        asin_dict = json.load(f)
                        all_asin_data.append(asin_dict)
                        # Dynamically expand columns list sequentially if a brand new feature shows up
                        for key in asin_dict.keys():
                            if key not in global_headers:
                                global_headers.append(key)

    # 2. Build the master Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Products Scraped Data"
    ws.views.sheetView[0].showGridLines = True

    # Append global headers row
    ws.append(global_headers)

    # Append structured dataset rows dynamically matching elements into their precise columns
    for data_dict in all_asin_data:
        row_values = []
        for header in global_headers:
            # Map values matching key, else leave structural cell blank
            row_values.append(data_dict.get(header, ""))
        ws.append(row_values)

    # 3. Apply Professional Styling Layout
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    body_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'), right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'), bottom=Side(style='thin', color='E5E7EB')
    )

    # Format Header Row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Format All Data Rows
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Set uniform comfortable width parameters across dynamically parsed items
    for col in ws.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 30

    # Save Master File directly into base zip root session folder
    master_excel_filepath = os.path.join(target_session_path, "Master_Product_Details.xlsx")
    wb.save(master_excel_filepath)

    # --- ZIP ENGINE COMPRESSION ROUTINE ---
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for root, dirs, files in os.walk(target_session_path):
            for file in files:
                # Exclude internal asset storage JSON components from finalized output zip presentation
                if file == "data.json":
                    continue
                full_path = os.path.join(root, file)
                relative_path = os.path.relpath(full_path, target_session_path)
                zip_file.write(full_path, relative_path)

    zip_buffer.seek(0)

    try:
        shutil.rmtree(target_session_path)
    except Exception: pass

    return send_file(zip_buffer, mimetype="application/zip", as_attachment=True, download_name="Amazon_Bulk_Scraped_Data.zip")

HOST = "0.0.0.0"  # bind every interface: reachable on localhost AND this machine's LAN IP
PORT = 5000

def get_lan_ip():
    """ Best-effort LAN address, just so startup can print a usable URL.

    Opening a UDP socket toward a public address makes Windows pick the outbound
    interface; nothing is actually transmitted.
    """
    try:
        probe = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            probe.connect(("8.8.8.8", 80))
            return probe.getsockname()[0]
        finally:
            probe.close()
    except Exception:
        return None

if __name__ == "__main__":
    lan_ip = get_lan_ip()
    print(" * Amazon Scraper is reachable at:")
    print(f"     http://localhost:{PORT}")
    print(f"     http://127.0.0.1:{PORT}")
    if lan_ip:
        print(f"     http://{lan_ip}:{PORT}   <- this machine on the network")
    print(f" * Swagger UI: /docs     OpenAPI spec: /openapi.json")
    print(" * Serving on every interface. Windows Firewall must allow inbound TCP", PORT)

    app.run(host=HOST, debug=False, use_reloader=False, port=PORT)

# if __name__ == "__main__":
#     parser = argparse.ArgumentParser()
#     parser.add_argument("--port", type=int, default=5000)
#     args = parser.parse_args()

#     app.run(
#         debug=False,
#         use_reloader=False,
#         port=args.port
#     )